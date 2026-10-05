#!/usr/bin/env python3
"""Автотесты docs2md: чистка Markdown и конвертация всех поддержанных форматов.

Что проверяем и почему именно так:
  1) **чистка Markdown** — она снимает артефакты экранирования, но обязана НЕ трогать код, таблицы
     и идентификаторы вида `__pycache__`; такие вещи ломаются при первой же правке, поэтому закреплены;
  2) **конвертация форматов** — по файлу каждого поддержанного формата: так ловится поломка
     зависимостей. Именно этот прогон однажды поймал `.ipynb`, который валился из-за отсутствующих
     шаблонов nbconvert.

Запуск:  python3 tests/test_docs2md.py      (код 0 — всё хорошо, 1 — есть расхождения)
"""
from __future__ import annotations

import json
import os
import sys
import tempfile
import zipfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "scripts"))

import docs2md          # noqa: E402  — проверяем именно точку входа
import doc_to_md        # noqa: E402

CASES: list[tuple[str, str, str]] = [
    ("снятие экранирования", r"07\.09\.2026 \(выдержка\)", "07.09.2026 (выдержка)"),
    ("жирный фразой", "__сервисный дом бизнес-класса__", "**сервисный дом бизнес-класса**"),
    ("жирный одним словом по-русски", "__ХИМКИ__", "**ХИМКИ**"),
    ("дундеры не трогаем", "папка __pycache__ и метод __init__", "папка __pycache__ и метод __init__"),
    ("snake_case не трогаем", "поле snake__case", "поле snake__case"),
    ("блок кода не трогаем", "```\na\\.b\n```", "```\na\\.b\n```"),
    ("inline-код не трогаем", "`x\\.y`", "`x\\.y`"),
    ("таблица: экранирование «|» остаётся", "| a\\|b |\n|---|\n| c\\|d |", "| a\\|b |\n|---|\n| c\\|d |"),
    ("путь Windows не ломаем", r"C:\temp\new", r"C:\temp\new"),
    ("англ. одиночное жирное оставляем", "__Summary__", "__Summary__"),
]


def _minimal_docx(path: str) -> None:
    """Крошечный .docx руками — чтобы проверить mammoth без образца в репозитории."""
    doc = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>
<w:p><w:r><w:t>Проверка DOCX 07.09.2026</w:t></w:r></w:p>
<w:p><w:r><w:rPr><w:b/></w:rPr><w:t>жирная строка</w:t></w:r></w:p>
</w:body></w:document>"""
    types = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>"""
    rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>"""
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", types)
        z.writestr("_rels/.rels", rels)
        z.writestr("word/document.xml", doc)


def _sample_pdf(path: str) -> None:
    """PDF с текстом — проверяем, что текст и заголовок страницы попадают в Markdown."""
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 100), "Отчёт по рынку LiFePO4, 2026", fontsize=14)
    page.insert_text((72, 130), "Проверка извлечения текста из PDF.", fontsize=11)
    doc.save(path)
    doc.close()


def _sample_pptx(path: str) -> None:
    """PPTX с текстом и таблицей — проверяем оба пути разбора."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(6), Inches(1))
    box.text_frame.text = "Слайд про TPMS"

    rows, cols = 2, 2
    table = slide.shapes.add_table(rows, cols, Inches(0.5), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "колесо"
    table.cell(0, 1).text = "давление"
    table.cell(1, 0).text = "переднее"
    table.cell(1, 1).text = "2.2"
    prs.save(path)


def _build_samples(d: str) -> list[tuple[str, int]]:
    """По файлу каждого поддержанного формата. Возвращает (имя, минимум знаков)."""
    from openpyxl import Workbook
    from ebooklib import epub

    _minimal_docx(os.path.join(d, "t.docx"))
    _sample_pdf(os.path.join(d, "t.pdf"))
    _sample_pptx(os.path.join(d, "t.pptx"))

    wb = Workbook()
    ws = wb.active
    ws.title = "Лист1"
    ws.append(["статья", "сумма"])
    ws.append(["аренда", 120000])
    wb.save(os.path.join(d, "t.xlsx"))
    wb.save(os.path.join(d, "t.xlsm"))

    nb = {"cells": [
        {"cell_type": "markdown", "metadata": {}, "source": ["# Заголовок\n", "текст"]},
        {"cell_type": "code", "metadata": {}, "execution_count": 1,
         "outputs": [{"output_type": "stream", "name": "stdout", "text": ["вывод\n"]}],
         "source": ["print(1)"]},
    ], "metadata": {"language_info": {"name": "python"}}, "nbformat": 4, "nbformat_minor": 5}
    with open(os.path.join(d, "t.ipynb"), "w", encoding="utf-8") as fh:
        fh.write(json.dumps(nb))

    book = epub.EpubBook()
    book.set_identifier("id1")
    book.set_title("Тест")
    ch = epub.EpubHtml(title="Глава", file_name="ch1.xhtml", lang="ru")
    ch.content = "<h1>Привет</h1><p>Ёж.</p>"
    book.add_item(ch)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav", ch]
    epub.write_epub(os.path.join(d, "t.epub"), book)

    with open(os.path.join(d, "t.html"), "w", encoding="utf-8") as fh:
        fh.write("<h1>Заголовок</h1><table><tr><td>a</td><td>b</td></tr></table>")
    for name, body in (("t.txt", "просто текст\n"), ("t.md", "# заголовок\n"),
                       ("t.csv", "a,b\n1,2\n"), ("t.json", '{"a":1}'), ("t.yaml", "a: 1\n")):
        with open(os.path.join(d, name), "w", encoding="utf-8") as fh:
            fh.write(body)

    return [("t.docx", 10), ("t.pdf", 20), ("t.pptx", 20), ("t.html", 10), ("t.epub", 10),
            ("t.ipynb", 10), ("t.xlsx", 10), ("t.xlsm", 10), ("t.txt", 5), ("t.md", 5),
            ("t.csv", 3), ("t.json", 3), ("t.yaml", 3)]


def test_cleanup() -> int:
    bad = 0
    print("— чистка Markdown:")
    for name, src, want in CASES:
        got = doc_to_md.clean_markdown(src)
        if got == want:
            print(f"  ✅ {name}")
        else:
            bad += 1
            print(f"  ❌ {name}\n     было:  {src!r}\n     стало: {got!r}\n     ждали: {want!r}")
    return bad


def test_formats() -> int:
    bad = 0
    print("— конвертация форматов (через docs2md.convert):")
    with tempfile.TemporaryDirectory() as d:
        for name, min_chars in _build_samples(d):
            path = os.path.join(d, name)
            try:
                out = docs2md.convert(path)
            except SystemExit as exc:
                bad += 1
                print(f"  ❌ {name}: отказ — {exc}"); continue
            except Exception as exc:                       # noqa: BLE001
                bad += 1
                print(f"  ❌ {name}: {type(exc).__name__}: {exc}"); continue
            if len(out.strip()) >= min_chars:
                print(f"  ✅ {name}: {len(out)} знаков")
            else:
                bad += 1
                print(f"  ❌ {name}: подозрительно пусто ({len(out)} знаков)")
    return bad


def test_details() -> int:
    """Точечные проверки: таблица PPTX, жирный DOCX, отказ на неподдержанном формате."""
    bad = 0
    print("— детали:")
    with tempfile.TemporaryDirectory() as d:
        p = os.path.join(d, "t.pptx")
        _sample_pptx(p)
        out = docs2md.convert(p)
        for needle, label in (("Слайд про TPMS", "текст слайда"), ("| колесо | давление |", "таблица слайда")):
            if needle in out:
                print(f"  ✅ PPTX: {label} извлечён")
            else:
                bad += 1
                print(f"  ❌ PPTX: {label} потерян")

        p = os.path.join(d, "t.docx")
        _minimal_docx(p)
        out = docs2md.convert(p)
        if "**жирная строка**" in out:
            print("  ✅ DOCX: жирный текст стал **жирным**, экранирования нет")
        else:
            bad += 1
            print(f"  ❌ DOCX: неожиданный вывод: {out[:120]!r}")

        try:
            with open(os.path.join(d, "t.doc"), "w") as fh:
                fh.write("x")
            docs2md.convert(os.path.join(d, "t.doc"))
            bad += 1
            print("  ❌ .doc должен был дать честный отказ")
        except SystemExit:
            print("  ✅ .doc: честный отказ вместо «почти сработало»")
    return bad


def main() -> int:
    bad = test_cleanup() + test_formats() + test_details()
    print(f"\nитого провалено: {bad}")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
