#!/usr/bin/env python3
"""PPTX → Markdown: текст, таблицы, заметки слайдов.

Зачем отдельный скрипт: навык умел презентации только СОЗДАВАТЬ, а читать чужой PPTX — нет.

Важная деталь (проверено на живых презентациях): текст в PPTX часто лежит **внутри групп фигур**.
Без рекурсии в группы теряется почти всё содержимое — на презентации из 20 слайдов остаётся меньше
тысячи знаков вместо шестнадцати тысяч. Поэтому `_shape_md` заходит в группы.

Использование:
    python3 pptx_to_md.py презентация.pptx > презентация.md
"""
from __future__ import annotations

import sys

PICTURE = 13          # MSO_SHAPE_TYPE.PICTURE


def _shape_md(shape, out: list[str]) -> None:
    """Разобрать одну фигуру: группа → рекурсия, таблица → Markdown, текст → абзацы."""
    stype = shape.shape_type

    if stype is not None and str(stype).startswith("GROUP"):
        for sub in shape.shapes:
            _shape_md(sub, out)
        return

    if getattr(shape, "has_table", False):
        rows = [[c.text.replace("\n", " ").strip() for c in r.cells] for r in shape.table.rows]
        if rows:
            width = max(len(r) for r in rows)
            for i, row in enumerate(rows):
                cells = row + [""] * (width - len(row))
                out.append("| " + " | ".join(cells) + " |")
                if i == 0:
                    out.append("|" + "---|" * width)
            out.append("")
        return

    if stype == PICTURE:
        out.append(f"_[картинка: {shape.name}]_")
        return

    if getattr(shape, "has_text_frame", False):
        for p in shape.text_frame.paragraphs:
            text = "".join(r.text for r in p.runs).strip()
            if not text:
                continue
            level = p.level or 0
            out.append(("  " * level) + ("- " if level else "") + text)


def to_markdown(path: str) -> str:
    """Превратить .pptx в Markdown. Заметки слайдов добавляются цитатой."""
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation(path)
    out: list[str] = [f"# {path.split('/')[-1]}", ""]
    out.append(f"_слайдов: {len(prs.slides)}; размер: "
               f"{Emu(prs.slide_width).inches:.1f}×{Emu(prs.slide_height).inches:.1f} in_")
    out.append("")

    for number, slide in enumerate(prs.slides, 1):
        body: list[str] = []
        for shape in slide.shapes:
            _shape_md(shape, body)
        notes = ""
        if slide.has_notes_slide:
            notes = (slide.notes_slide.notes_text_frame.text or "").strip()
        out.append(f"## Слайд {number}")
        out.append("\n".join(body) if body else "_(текста нет)_")
        if notes:
            out.append(f"\n> заметки: {notes[:300]}")
        out.append("")

    return "\n".join(out).strip() + "\n"


def main() -> int:
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    print(to_markdown(sys.argv[1]), end="")
    return 0


if __name__ == "__main__":
    sys.exit(main())
