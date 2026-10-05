#!/usr/bin/env python3
"""Документ → Markdown. Поддержка: .docx, .html/.htm, .epub, .ipynb, .xlsx/.xlsm, .txt/.md/.csv.

Старые форматы (.doc, .odt, .rtf, .tex, .rst, .org, .typ) требуют pandoc — его нет,
поэтому честно сообщаем об этом, а не пытаемся «как-нибудь».

Примеры:
    doc_to_md.py "договор.docx" > договор.md
    doc_to_md.py "статья.html" --out статья.md
    doc_to_md.py "книга.epub"
"""
from __future__ import annotations

import argparse
import os
import re
import sys

LEGACY_HINT = ("нужен pandoc (не установлен). Это форматы .doc/.odt/.rtf/.tex/.rst/.org/.typ — "
               "скажите, если нужны: установим отдельно.")

# ── Чистка артефактов экранирования (mammoth и markdownify экранируют знаки препинания)
# Проверено 05.10.2026: mammoth выдавал «07\.09\.2026», «\(выдержка\)», «__жирный__» —
# в тексте это мусор, который мешает и человеку, и поиску по документу.
#
# ВАЖНО: чистка НЕ трогает блоки кода, inline-код и вертикальную черту.
#   * Код: трогать нельзя в принципе.
#   * `\|` — это правильное экранирование внутри таблиц Markdown; сняв его, мы бы сломали таблицы.
_ESCAPE_RE = re.compile(r"\\([.()\-_#*\[\]!+~<>{}$])")
_BOLD_RE = re.compile(r"(?<![\w_])__([^_\n]+?)__(?![\w_])")
_CODE_FENCE_RE = re.compile(r"(```.*?```|~~~.*?~~~)", re.S)
_INLINE_CODE_RE = re.compile(r"(`+[^`]*`+)")


def _clean_chunk(s: str) -> str:
    s = _ESCAPE_RE.sub(r"\1", s)
    s = _BOLD_RE.sub(_bold_fix, s)
    return s


def _bold_fix(m: "re.Match[str]") -> str:
    """`__жирный__` → `**жирный**`, но `__pycache__` и `__init__` не трогаем.

    Различить их регуляркой нельзя, поэтому смотрим на содержимое: у настоящего жирного
    (его ставит mammoth) почти всегда есть пробел или кириллица. Если внутри только ASCII
    без пробелов — это похоже на идентификатор, и трогать его нельзя: `**pycache**` в
    документации про Python — это ошибка. Оставленный `__Summary__` при этом всё равно
    отрисуется жирным, так что худший случай безобиден.
    """
    inner = m.group(1)
    if " " in inner or any(ord(ch) > 127 for ch in inner):
        return f"**{inner}**"
    return m.group(0)


def clean_markdown(text: str) -> str:
    """Снять артефакты экранирования, не задев код и таблицы."""
    parts = _CODE_FENCE_RE.split(text)          # нечётные индексы — блоки кода
    out: list[str] = []
    for i, part in enumerate(parts):
        if i % 2:
            out.append(part)
            continue
        sub = _INLINE_CODE_RE.split(part)       # нечётные индексы — inline-код
        out.append("".join(x if j % 2 else _clean_chunk(x) for j, x in enumerate(sub)))
    return "".join(out)


def from_docx(path: str) -> str:
    import mammoth

    with open(path, "rb") as fh:
        result = mammoth.convert_to_markdown(fh)
    text = result.value
    if result.messages:
        notes = "; ".join(f"{m.type}: {m.message}" for m in result.messages[:5])
        text += f"\n\n<!-- предупреждения конвертера: {notes} -->\n"
    return text


def from_html(path: str) -> str:
    from markdownify import markdownify

    with open(path, encoding="utf-8", errors="replace") as fh:
        return markdownify(fh.read(), heading_style="ATX")


def from_epub(path: str) -> str:
    import ebooklib
    from ebooklib import epub
    from markdownify import markdownify

    book = epub.read_epub(path)
    parts: list[str] = []
    for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
        html = item.get_content().decode("utf-8", errors="replace")
        md = markdownify(html, heading_style="ATX").strip()
        if md:
            parts.append(md)
    return "\n\n---\n\n".join(parts)


def from_ipynb(path: str) -> str:
    """Свой конвертер .ipynb — без nbconvert.

    Почему не nbconvert: его MarkdownExporter ищет шаблоны Jupyter, а в нашей пользовательской
    установке их нет — падает с «No template sub-directory with name 'markdown'» (проверено
    05.10.2026). Формат .ipynb — обычный JSON со списком ячеек, поэтому читаем его сами:
    markdown-ячейки отдаём как есть, код — в блоке с языком, выводы — сжато и в блоках.
    """
    import json

    with open(path, encoding="utf-8", errors="replace") as fh:
        nb = json.load(fh)

    lang = str((((nb.get("metadata") or {}).get("language_info") or {}).get("name")) or "")
    out: list[str] = []

    def as_text(value) -> str:
        if isinstance(value, list):
            return "".join(str(x) for x in value)
        return str(value or "")

    for cell in nb.get("cells", []) or []:
        ctype = cell.get("cell_type")
        body = as_text(cell.get("source")).rstrip()

        if ctype == "markdown":
            out.append(body)
        elif ctype == "code":
            out.append(f"```{lang}\n{body}\n```")
            for o in cell.get("outputs") or []:
                otype = o.get("output_type")
                shown = ""
                if "text/plain" in (o.get("data") or {}):
                    shown = as_text(o["data"]["text/plain"]).rstrip()
                elif otype == "stream":
                    shown = as_text(o.get("text")).rstrip()
                elif otype == "error":
                    shown = f"{o.get('ename')}: {o.get('evalue')}"
                elif "image/png" in (o.get("data") or {}):
                    out.append("_[картинка из вывода]_")
                    continue
                if shown:
                    if len(shown) > 1000:
                        shown = shown[:1000] + "…"
                    out.append(f"```\n# вывод:\n{shown}\n```")
        elif ctype == "raw":
            out.append(body)
        out.append("")

    return "\n".join(out).strip() + "\n"


def from_xlsx(path: str) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    out: list[str] = []
    for ws in wb.worksheets:
        out.append(f"\n## Лист: {ws.title}  ({ws.max_row}×{ws.max_column})\n")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if c is None else str(c) for c in row]
            if not any(cells):
                continue
            out.append("| " + " | ".join(cells) + " |")
    wb.close()
    return "\n".join(out)


DISPATCH = {
    ".docx": from_docx,
    ".html": from_html,
    ".htm": from_html,
    ".epub": from_epub,
    ".ipynb": from_ipynb,
    ".xlsx": from_xlsx,
    ".xlsm": from_xlsx,
}
PLAIN = {".txt", ".md", ".markdown", ".csv", ".json", ".yaml", ".yml"}
LEGACY = {".doc", ".odt", ".rtf", ".tex", ".latex", ".rst", ".org", ".typ"}
# Форматы, где конвертер экранирует знаки препинания: чистим. У .ipynb и .xlsx экранирования нет,
# а у .txt/.md чистка вообще недопустима — файл должен остаться байт-в-байт.
CLEAN = {".docx", ".html", ".htm", ".epub"}


def convert(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext in DISPATCH:
        text = DISPATCH[ext](path)
        return clean_markdown(text) if ext in CLEAN else text
    if ext in PLAIN:
        with open(path, encoding="utf-8", errors="replace") as fh:
            return fh.read()
    if ext in LEGACY:
        raise SystemExit(f"{ext}: {LEGACY_HINT}")
    raise SystemExit(f"Не знаю формат {ext!r} — поддержаны: "
                     f"{', '.join(sorted(set(DISPATCH) | PLAIN))}")


def main() -> int:
    ap = argparse.ArgumentParser(description="Документ → Markdown")
    ap.add_argument("path")
    ap.add_argument("--out", help="записать в файл вместо stdout")
    args = ap.parse_args()

    if not os.path.isfile(args.path):
        print(f"Файл не найден: {args.path}", file=sys.stderr)
        return 1

    text = convert(args.path)
    if args.out:
        with open(args.out, "w", encoding="utf-8") as fh:
            fh.write(text)
        print(f"Записано: {args.out} ({len(text)} символов)", file=sys.stderr)
    else:
        print(text)
    return 0


if __name__ == "__main__":
    sys.exit(main())
